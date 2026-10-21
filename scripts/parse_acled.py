"""Llegeix els fitxers regionals agregats d'ACLED (raw/acled/*.xlsx).

Cada fitxer té una fila per setmana, regió administrativa (ADMIN1), tipus i subtipus
d'esdeveniment, amb EVENTS, FATALITIES i el centroide de la regió.

Sortida: build/acled.json
- adm: regions amb violència política (finestra de 12 i 52 setmanes)
- cw_fat / cw_ev: sèries setmanals de morts i esdeveniments per país (52 setmanes)

Si raw/acled/ és buida, no fa res i el procés continua amb el darrer build/acled.json.
"""
import glob, json, os, sys
from datetime import date, datetime, timedelta

ROOT = os.path.join(os.path.dirname(__file__), '..')
RAW = os.path.join(ROOT, 'raw', 'acled')
OUT = os.path.join(ROOT, 'build', 'acled.json')
PV = {'Battles', 'Explosions/Remote violence', 'Violence against civilians'}
FINESTRA_FOCUS = 12   # setmanes de la capa de focus del mapa
FINESTRA_PAIS = 52    # setmanes de les xifres per país
EXCEL0 = date(1899, 12, 30)


def to_date(v):
    if isinstance(v, datetime): return v.date()
    if isinstance(v, date): return v
    if isinstance(v, (int, float)): return EXCEL0 + timedelta(days=int(v))
    return datetime.strptime(str(v)[:10], '%Y-%m-%d').date()


def read(path):
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.worksheets[0]
    it = ws.iter_rows(values_only=True)
    head = [str(h).strip().upper() if h is not None else '' for h in next(it)]
    need = ['WEEK', 'COUNTRY', 'ADMIN1', 'EVENT_TYPE', 'EVENTS', 'FATALITIES', 'CENTROID_LATITUDE', 'CENTROID_LONGITUDE']
    missing = [c for c in need if c not in head]
    if missing:
        raise SystemExit(f'{os.path.basename(path)}: falten les columnes {missing}')
    ix = {c: head.index(c) for c in need}
    for r in it:
        if r[ix['WEEK']] is None: continue
        yield {c: r[i] for c, i in ix.items()}


def main():
    files = sorted(glob.glob(os.path.join(RAW, '*.xlsx')))
    # Si queden fitxers de mesos anteriors, només es fa servir el més recent de cada regió
    # (el nom acaba amb la data: ..._aggregated_data_up_to_week_of-AAAA-MM-DD.xlsx).
    per_regio = {}
    for f in files:
        regio = os.path.basename(f).split('_aggregated')[0]
        per_regio.setdefault(regio, []).append(f)
    files = [sorted(v)[-1] for v in per_regio.values()]
    for v in per_regio.values():
        for f in sorted(v)[:-1]:
            print(f'ACLED: s\'ignora {os.path.basename(f)} (n\'hi ha una versió més recent)', file=sys.stderr)
    if not files:
        print('ACLED: cap fitxer a raw/acled/; es manté el darrer build/acled.json', file=sys.stderr)
        return
    rows = []
    for f in files:
        n0 = len(rows); rows.extend(read(f))
        print(f'ACLED: {os.path.basename(f)} · {len(rows) - n0} files', file=sys.stderr)
    last = max(to_date(r['WEEK']) for r in rows)
    w52 = [last - timedelta(weeks=i) for i in range(FINESTRA_PAIS - 1, -1, -1)]
    first52, first12 = w52[0], last - timedelta(weeks=FINESTRA_FOCUS - 1)
    idx = {d: i for i, d in enumerate(w52)}
    adm, cw_f, cw_e = {}, {}, {}
    for r in rows:
        if r['EVENT_TYPE'] not in PV: continue
        d = to_date(r['WEEK'])
        if d < first52 or d not in idx: continue
        ev, fa = int(r['EVENTS'] or 0), int(r['FATALITIES'] or 0)
        c = r['COUNTRY']
        cw_f.setdefault(c, [0] * len(w52))[idx[d]] += fa
        cw_e.setdefault(c, [0] * len(w52))[idx[d]] += ev
        k = (c, r['ADMIN1'])
        a = adm.setdefault(k, {'c': c, 'a': r['ADMIN1'],
                               'll': [round(float(r['CENTROID_LONGITUDE']), 2), round(float(r['CENTROID_LATITUDE']), 2)],
                               'e12': 0, 'f12': 0, 'e52': 0, 'f52': 0, 't12': {}})
        a['e52'] += ev; a['f52'] += fa
        if d >= first12:
            a['e12'] += ev; a['f12'] += fa; a['t12'][r['EVENT_TYPE']] = a['t12'].get(r['EVENT_TYPE'], 0) + ev
    out = {'meta': {'source': 'ACLED aggregated data (regional files)', 'files': [os.path.basename(f) for f in files],
                    'weeks_from': first52.isoformat(), 'weeks_to': last.isoformat(), 'window12_from': first12.isoformat(),
                    'generated': date.today().isoformat(), 'event_types': sorted(PV)},
           'weeks': [d.isoformat() for d in w52],
           'adm': [a for a in adm.values() if a['f52'] > 0 or a['e12'] > 0],
           'cw_fat': cw_f, 'cw_ev': cw_e}
    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    json.dump(out, open(OUT, 'w'), ensure_ascii=False, separators=(',', ':'))
    print(f'ACLED: dades fins a la setmana del {last} · {len(out["adm"])} regions', file=sys.stderr)


if __name__ == '__main__':
    main()
